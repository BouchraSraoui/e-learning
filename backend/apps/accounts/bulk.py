from __future__ import annotations

import csv
import io

from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import transaction

from .models import Department, Language, Role, Team, User

COLUMNS = [
    'email', 'first_name', 'last_name', 'role', 'department', 'team',
    'manager', 'job_title', 'phone', 'location', 'language', 'is_active',
]

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000

_ROLE_VALUES = {r.value for r in Role}
_LEGACY_ROLE_ALIASES = {'learner': 'user', 'trainer': 'manager'}
_LANG_VALUES = {l.value for l in Language}
_TRUE = {'true', '1', 'yes', 'y', 'oui', 'active', 'x'}
_FALSE = {'false', '0', 'no', 'n', 'non', 'inactive', ''}



def _user_row(user: User) -> list[str]:
    return [
        user.email,
        user.first_name,
        user.last_name,
        user.role,
        user.department.name if user.department_id else '',
        user.team.name if user.team_id else '',
        user.manager.email if user.manager_id else '',
        user.job_title,
        user.phone,
        user.location,
        user.language,
        'true' if user.is_active else 'false',
    ]


def _export_csv(rows: list[list[str]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(COLUMNS)
    writer.writerows(rows)
    return buf.getvalue().encode('utf-8-sig')


def _export_xlsx(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = 'Users'
    ws.append(COLUMNS)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def build_export(queryset, fmt: str) -> tuple[bytes, str, str]:
    rows = [_user_row(u) for u in queryset]
    if fmt == 'xlsx':
        return _export_xlsx(rows), _XLSX_TYPE, 'users.xlsx'
    return _export_csv(rows), 'text/csv; charset=utf-8', 'users.csv'


def build_template(fmt: str) -> tuple[bytes, str, str]:
    example = [
        'new.user@gmail.com', 'Amina', 'Benali', 'user', 'Commercial',
        '', 'manager@gmail.com', 'Account Manager', '', 'Algiers', 'fr', 'true',
    ]
    if fmt == 'xlsx':
        return _export_xlsx([example]), _XLSX_TYPE, 'users-template.xlsx'
    return _export_csv([example]), 'text/csv; charset=utf-8', 'users-template.csv'



def _norm_header(name) -> str:
    return str(name or '').strip().lower().replace(' ', '_')


def _cell(value) -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'true' if value else 'false'
    return str(value).strip()


def _read_rows(upload) -> list[dict]:
    name = (getattr(upload, 'name', '') or '').lower()
    if name.endswith('.xlsx'):
        return _read_xlsx(upload)
    if name.endswith('.csv') or not name:
        return _read_csv(upload)
    raise ValueError('Unsupported file type — upload a .csv or .xlsx file.')


def _read_csv(upload) -> list[dict]:
    raw = upload.read()
    if isinstance(raw, bytes):
        text = raw.decode('utf-8-sig', errors='replace')
    else:
        text = raw
    reader = csv.reader(io.StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        return []
    keys = [_norm_header(h) for h in header]
    rows = []
    for values in reader:
        rows.append({k: _cell(v) for k, v in zip(keys, values)})
    return rows


def _read_xlsx(upload) -> list[dict]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError('Could not read the XLSX file.') from exc
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        return []
    keys = [_norm_header(h) for h in header]
    rows = []
    for values in it:
        if values is None or all(v is None or _cell(v) == '' for v in values):
            continue
        rows.append({k: _cell(v) for k, v in zip(keys, values)})
    return rows


def _resolve_names(rows: list[dict]) -> tuple[dict, dict, dict]:
    departments = {d.name.lower(): d for d in Department.objects.all()}
    teams = {t.name.lower(): t for t in Team.objects.select_related('department')}
    emails = {r.get('manager', '').strip().lower() for r in rows if r.get('manager')}
    emails |= {r.get('email', '').strip().lower() for r in rows if r.get('email')}
    users = {u.email.lower(): u for u in User.objects.filter(email__in=emails)}
    return departments, teams, users


def _apply_row(row: dict, *, departments, teams, existing_users) -> tuple[str, list[str]]:
    errors: list[str] = []

    email = row.get('email', '').strip().lower()
    if not email:
        return 'error', ['Missing email.']
    try:
        validate_email(email)
    except DjangoValidationError:
        return 'error', [f'Invalid email: {email!r}.']

    user = existing_users.get(email)
    is_new = user is None

    first = row.get('first_name', '').strip()
    last = row.get('last_name', '').strip()
    full = row.get('full_name', '').strip()
    if full and not (first or last):
        parts = full.split()
        first = parts[0] if parts else full
        last = ' '.join(parts[1:])
    if is_new and not first:
        errors.append('Missing first_name (or full_name).')

    role = row.get('role', '').strip().lower()
    role = _LEGACY_ROLE_ALIASES.get(role, role)
    if role and role not in _ROLE_VALUES:
        errors.append(f'Invalid role: {role!r}.')
        role = ''

    language = row.get('language', '').strip().lower()
    if language and language not in _LANG_VALUES:
        errors.append(f'Invalid language: {language!r}.')
        language = ''

    department = None
    dep_name = row.get('department', '').strip()
    if dep_name:
        department = departments.get(dep_name.lower())
        if department is None:
            errors.append(f'Unknown department: {dep_name!r}.')

    team = None
    team_name = row.get('team', '').strip()
    if team_name:
        team = teams.get(team_name.lower())
        if team is None:
            errors.append(f'Unknown team: {team_name!r}.')
        elif department and team.department_id and team.department_id != department.id:
            errors.append(f'Team {team_name!r} is not in department {dep_name!r}.')

    manager = None
    mgr_email = row.get('manager', '').strip().lower()
    if mgr_email:
        manager = existing_users.get(mgr_email)
        if manager is None:
            errors.append(f'Unknown manager email: {mgr_email!r}.')
        elif not is_new and user and manager.pk == user.pk:
            errors.append('A user cannot be their own manager.')

    is_active = True
    raw_active = row.get('is_active', '').strip().lower()
    if raw_active:
        if raw_active in _TRUE:
            is_active = True
        elif raw_active in _FALSE:
            is_active = False
        else:
            errors.append(f'Invalid is_active: {raw_active!r}.')

    password = row.get('password', '').strip()
    if password:
        try:
            validate_password(password)
        except DjangoValidationError as exc:
            errors.append(f'Weak password: {"; ".join(exc.messages)}')

    if errors:
        return 'error', errors

    with transaction.atomic():
        if is_new:
            user = User(email=email)
        if first or is_new:
            user.first_name = first
        if last or is_new:
            user.last_name = last
        if role:
            user.role = role
        elif is_new:
            user.role = Role.USER
        if language:
            user.language = language
        if dep_name:
            user.department = department
        if team_name:
            user.team = team
        if mgr_email:
            user.manager = manager
        for field in ('job_title', 'phone', 'location'):
            val = row.get(field, '').strip()
            if val:
                setattr(user, field, val)
        if raw_active or is_new:
            user.is_active = is_active
        if password:
            user.set_password(password)
        elif is_new:
            user.set_unusable_password()
        user.save()
        existing_users[email] = user

    return ('created' if is_new else 'updated'), []


def run_import(upload, *, actor=None) -> dict:
    size = getattr(upload, 'size', None)
    if size and size > MAX_IMPORT_BYTES:
        raise ValueError('File is too large (max 5 MB).')

    rows = _read_rows(upload)
    if len(rows) > MAX_IMPORT_ROWS:
        raise ValueError(f'Too many rows (max {MAX_IMPORT_ROWS}).')

    departments, teams, existing = _resolve_names(rows)

    report = {'total': len(rows), 'created': 0, 'updated': 0, 'errors': []}
    for index, row in enumerate(rows, start=2):
        try:
            outcome, errors = _apply_row(
                row, departments=departments, teams=teams, existing_users=existing
            )
        except Exception as exc:
            outcome, errors = 'error', [f'Unexpected error: {exc}']
        if outcome == 'error':
            report['errors'].append(
                {'row': index, 'email': row.get('email', '').strip(), 'messages': errors}
            )
        else:
            report[outcome] += 1
    return report
